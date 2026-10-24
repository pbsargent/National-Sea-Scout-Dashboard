"""Read only reviewed aggregate cells from the captured membership history tab."""
from math import isfinite
from openpyxl import load_workbook

HISTORY_SHEET = '2010-202X Membership'
FIELDS = {
    'Councils with Ships': 'councils', '# Ships': 'ships',
    'Primary Youth': 'primaryYouth', 'Multiple Youth': 'multipleYouth',
    'Total Youth': 'youth', 'primary registered adults': 'primaryAdults',
    'total registered adults': 'adults',
}


def count(value):
    if isinstance(value, bool) or not isinstance(value, (float, int)):
        return None
    return int(value) if isfinite(value) and value >= 0 and value == int(value) else None


def read_history(capture_path):
    workbook = load_workbook(capture_path, read_only=True, data_only=True)
    try:
        if HISTORY_SHEET not in workbook.sheetnames:
            return {'sourceSheet': HISTORY_SHEET, 'annual': [], 'note': 'History worksheet unavailable in this capture.'}
        iterator = workbook[HISTORY_SHEET].iter_rows(values_only=True)
        headers = next(iterator, ())
        columns = [(i, count(v)) for i, v in enumerate(headers) if count(v) is not None and 2000 <= count(v) <= 2100]
        records = {year: {'year': year, 'date': f'{year}-12-31', **{key: None for key in FIELDS.values()}} for _, year in columns}
        seen = set()
        for row in iterator:
            key = FIELDS.get(str(row[0] or '').strip()) if row else None
            if key and key not in seen:
                seen.add(key)
                for i, year in columns:
                    records[year][key] = count(row[i]) if i < len(row) else None
        return {'sourceSheet': HISTORY_SHEET, 'annual': [records[y] for y in sorted(records)],
                'note': 'Annual columns as recorded in the captured workbook. December 31 identifies the annual comparison period; source revision dates are unavailable. These figures may differ from historical slide snapshots.'}
    finally:
        workbook.close()
