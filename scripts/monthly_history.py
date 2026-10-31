"""Read allowlisted national Ship monthly observations from TY History."""
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook
from membership_history import count


def read_monthly_history(path):
    result = {'sourceSheet': 'TY History', 'sourceCaptureName': Path(path).name, 'observations': []}
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if 'TY History' not in book.sheetnames:
            return result
        rows = book['TY History'].iter_rows(values_only=True)
        headers = list(next(rows, ()))
        required = ['Date', 'Level_Type', 'Level_Number', 'Unit_Type', 'Total_Units', 'Total_Adults', 'Total_Youth']
        if not all(name in headers for name in required):
            raise ValueError('TY History is missing required columns')
        seen = set()
        for values in rows:
            row = dict(zip(headers, values))
            if (row.get('Level_Type'), row.get('Level_Number'), row.get('Unit_Type')) != ('National', 0, 'Ship'):
                continue
            day = row.get('Date')
            if isinstance(day, datetime):
                day = day.date()
            if not isinstance(day, date):
                raise ValueError('TY History has an invalid observation date')
            month = day.strftime('%Y-%m')
            if month in seen:
                raise ValueError(f'TY History has duplicate national Ship month: {month}')
            seen.add(month)
            result['observations'].append({'date': day.isoformat(), 'ships': count(row.get('Total_Units')),
                                           'adults': count(row.get('Total_Adults')), 'youth': count(row.get('Total_Youth'))})
        result['observations'].sort(key=lambda row: row['date'])
        return result
    finally:
        book.close()
